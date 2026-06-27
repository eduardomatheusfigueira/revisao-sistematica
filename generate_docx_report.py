import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def add_heading_styled(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.bold = True
    
    if level == 1:
        run.font.size = Pt(14)
    else:
        run.font.size = Pt(12)
    return p

def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.49) # ~1.25 cm
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p

def main():
    doc = docx.Document()
    
    # Configure A4 page margins (2.5 cm all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)  # A4 width
        section.page_height = Inches(11.69) # A4 height
        
    # --- Title Page / Header ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(12)
    title_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("Metodologia e Operação do Sistema de Coleta Multibases para Revisões Sistemáticas")
    title_run.font.name = 'Times New Roman'
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    
    author_p = doc.add_paragraph()
    author_p.paragraph_format.space_after = Pt(24)
    author_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_run = author_p.add_run("Pair Programming System Development - Technical Article Report")
    author_run.font.name = 'Times New Roman'
    author_run.font.size = Pt(12)
    author_run.font.italic = True
    
    # --- Abstract ---
    abstract_heading = doc.add_paragraph()
    abstract_heading.paragraph_format.space_after = Pt(6)
    abstract_heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ab_run = abstract_heading.add_run("RESUMO")
    ab_run.font.name = 'Times New Roman'
    ab_run.font.size = Pt(12)
    ab_run.font.bold = True
    
    p_abstract = doc.add_paragraph()
    p_abstract.paragraph_format.space_after = Pt(24)
    p_abstract.paragraph_format.line_spacing = 1.15
    p_abstract.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abstract.paragraph_format.left_indent = Inches(0.5)
    p_abstract.paragraph_format.right_indent = Inches(0.5)
    ab_text_run = p_abstract.add_run(
        "Este artigo descreve a concepção, o desenvolvimento e a validação do sistema integrado de desenho de pesquisa, coleta e gestão de literatura científica para revisões sistemáticas. O ecossistema é guiado por um formulário de planejamento em Excel integrado a coletores independentes dedicados a quatro bases: a Biblioteca Digital Brasileira de Teses e Dissertações (BDTD), a Scientific Electronic Library Online (SciELO), a Scopus da Elsevier e o diretório aberto OpenAlex. A arquitetura emprega consultas via APIs REST e técnicas avançadas de web scraping em linguagem Python, com um algoritmo híbrido de enriquecimento de resumos. Como diferencial, o sistema unifica o protocolo e as planilhas de gestão operacionais (triagem, extração de dados, risco de viés e PRISMA) em um único ambiente Excel auto-exportável, garantindo reprodutibilidade, padronização e eficiência em todo o ciclo de vida da revisão."
    )
    ab_text_run.font.name = 'Times New Roman'
    ab_text_run.font.size = Pt(11)
    
    # --- 1. Introdução ---
    add_heading_styled(doc, "1. Introdução", 1)
    
    add_body_paragraph(doc, 
        "A realização de revisões sistemáticas da literatura acadêmica constitui a base das pesquisas científicas rigorosas, exigindo um mapeamento exaustivo da produção científica. O processo tradicional de levantamento manual nas diferentes bases de dados apresenta-se lento, passível de erros humanos e extremamente fragmentado. Cada base científica opera com padrões de busca, formatos de indexação e limitações de APIs distintos, o que dificulta a consolidação rápida de dados bibliométricos. Visando contornar esses obstáculos, o presente projeto implementa uma suíte de coletores automáticos unificados em linguagem Python."
    )
    
    add_body_paragraph(doc,
        "O sistema automatizado aborda a coleta a partir de quatro fontes principais com finalidades complementares para a pesquisa científica nacional e internacional. A Biblioteca Digital Brasileira de Teses e Dissertações (BDTD) fornece o acesso direto à produção acadêmica de pós-graduação stricto sensu nacional. A SciELO atua como o principal diretório de periódicos de acesso aberto da América Latina e Caribe. Por sua vez, a base comercial Scopus da Elsevier oferece indexação global de revistas de alto impacto, enquanto o catálogo aberto OpenAlex opera como um agregador unificado de toda a literatura mundial."
    )
    
    add_body_paragraph(doc,
        "Além da automatização das buscas textuais simples, a suite resolve problemas comuns de padronização estrutural de metadados e restrição de acesso a resumos. Coletores acadêmicos frequentemente esbarram em firewalls de segurança agressivos ou em limitações contratuais que ocultam o texto completo dos resumos. A arquitetura aqui proposta utiliza técnicas combinadas de web scraping polido e enriquecimento cruzado de dados entre diferentes APIs. Este artigo detalha o funcionamento metodológico de cada coletor e apresenta as instruções para a operação local do software."
    )
    
    # --- 2. Arquitetura do Sistema ---
    add_heading_styled(doc, "2. Arquitetura Geral do Sistema", 1)
    
    add_body_paragraph(doc,
        "A suíte de ferramentas foi projetada utilizando o paradigma de scripts modulares independentes e orientados por arquivos de configuração estruturados no formato JSON. Cada coletor opera de forma isolada, contendo seu próprio inicializador em batch (arquivos .bat) e rotinas internas de validação. Essa modularidade assegura que modificações nos layouts de uma determinada plataforma não causem impactos na operação das demais. O núcleo de processamento utiliza bibliotecas consolidadas do ecossistema Python, tais como Requests para chamadas HTTP e Beautiful Soup para parseamento de HTML."
    )
    
    add_body_paragraph(doc,
        "O fluxo de armazenamento de dados foi padronizado em todas as ferramentas para assegurar a consistência dos metadados exportados ao final do processo. Primeiramente, o sistema grava os resultados brutos da API em arquivos no formato JSON, servindo como uma cópia de segurança para auditorias posteriores. Em seguida, os dados são limpos de tags HTML e consolidados em planilhas tabulares CSV com quatorze colunas predefinidas. Por fim, uma planilha Excel com formatação rica de cores corporativas (identidade visual de cada base) e abas de sumários estatísticos é gerada."
    )
    
    add_body_paragraph(doc,
        "Adicionalmente, os coletores criam automaticamente relatórios markdown estatísticos para oferecer um diagnóstico imediato sobre o levantamento bibliométrico realizado. Estes relatórios Markdown exibem a cronologia das publicações coletadas, a distribuição dos idiomas de resumo, o ranking dos periódicos mais frequentes e a tabela dos artigos mais citados. Com isso, os pesquisadores conseguem avaliar a relevância do conjunto de dados coletado sem a necessidade de importar as tabelas em softwares externos. Logs detalhados registram todo o fluxo e erros de execução em arquivos específicos."
    )
    
    # --- 3. Coleta na BDTD ---
    add_heading_styled(doc, "3. Coleta na Biblioteca Brasileira de Teses e Dissertações", 1)
    
    add_body_paragraph(doc,
        "O coletor da BDTD conecta-se diretamente à API oficial de busca mantida pelo Instituto Brasileiro de Informação em Ciência e Tecnologia (IBICT). O script foi construído para realizar a paginação de registros de forma flexível utilizando parâmetros de paginação sequenciais nativos da plataforma. A busca recupera dados de autoria, instituição de defesa, grau acadêmico (mestrado ou doutorado), ano de publicação e o link direto para download do documento original. A rotina foi adaptada para manter o rigor metodológico exigido em investigações sistemáticas nacionais."
    )
    
    add_body_paragraph(doc,
        "A extração de resumos da BDTD lida com a inconsistência linguística comum nas teses e dissertações brasileiras, onde resumos podem estar em múltiplos idiomas. O script verifica se a API retornou o texto do resumo em português ou inglês e os concatena no campo final para evitar perdas textuais na análise conceitual. Além disso, as colunas do arquivo final foram adaptadas para refletir a terminologia da pós-graduação brasileira, em estrita conformidade com as diretrizes do Padrão de Metadados da BDTD (IBICT, 2017). O design visual das planilhas Excel da BDTD emprega tons de verde e amarelo, alinhados com a identidade nacional."
    )
    
    # --- 4. Coleta na Rede SciELO ---
    add_heading_styled(doc, "4. Coleta na Rede SciELO", 1)
    
    add_body_paragraph(doc,
        "O coletor da SciELO foi projetado utilizando a técnica de raspagem direta (web scraping) do portal oficial da SciELO (search.scielo.org) por meio de requisições estruturadas. Esta abordagem foi adotada após constatarmos que APIs secundárias abertas não possuíam o mapeamento exato da rede SciELO de periódicos latino-americanos e nacionais. O scraper utiliza paginação baseada no parâmetro de deslocamento nativo do portal, realizando requisições sequenciais de quinze em quinze resultados. Um parser BeautifulSoup lê a árvore do documento e mapeia os elementos estruturados, seguindo as melhores práticas de extração documentadas na literatura de engenharia de dados (MITCHELL, 2024)."
    )
    
    add_body_paragraph(doc,
        "Um diferencial deste coletor é o bypass estável do firewall e do WAF (Bunny Shield) implementado nos servidores da rede SciELO. O script simula perfeitamente os cabeçalhos de conexão de navegadores comerciais atualizados, mantendo cookies de sessão ativos e evitando bloqueios automáticos com erros HTTP 403. Os resumos multilíngues pré-carregados nas divs ocultas da própria página de resultados são mapeados pelo parser. Isso reduz a necessidade de requisições secundárias aos artigos individuais, otimizando o tempo de processamento da coleta substancialmente."
    )
    
    add_body_paragraph(doc,
        "O script da SciELO também conta com rotinas de fallback específicas para identificar e tratar de forma adequada preprints e artigos sem edições de periódicos definidos. Esses preprints são identificados pelo marcador visual da plataforma e seus anos de publicação são extraídos da string de data de disponibilização original. Desta forma, o script recupera exatamente a contagem total de artigos exibida no portal de busca da SciELO. As planilhas Excel deste coletor são estilizadas na tradicional identidade visual verde-floresta da marca SciELO."
    )
    
    # --- 5. Coleta na Base Scopus ---
    add_heading_styled(doc, "5. Coleta na Base Scopus", 1)
    
    add_body_paragraph(doc,
        "O coletor da Scopus consome diretamente a Scopus Search API fornecida oficialmente pela Elsevier. O script utiliza a chave de API (API Key) fornecida pelo usuário e configurada nas definições do sistema, enviada no cabeçalho X-ELS-APIKey. A busca suporta expressões complexas nativas da Scopus, como o filtro TITLE-ABS-KEY para correspondências no título, resumo ou palavras-chave dos artigos. A paginação utiliza o parâmetro de deslocamento inicial da busca e recupera informações completas de DOI, citações recebidas e dados bibliométricos oficiais."
    )
    
    add_body_paragraph(doc,
        "A API de busca da Scopus em seu plano básico restringe o acesso ao texto dos resumos e às palavras-chave completas para chaves sem credenciais institucionais autorizadas. Para garantir o funcionamento da coleta sob qualquer rede, o script valida e armazena os metadados bibliográficos oficiais recuperados da busca padrão. Caso o usuário execute o script conectado à rede VPN de uma instituição de pesquisa com assinatura ativa da Scopus, os resumos são coletados de forma nativa e automática pela API. A identidade visual adotada nos arquivos do coletor Scopus utiliza tons clássicos de azul-aço."
    )
    
    # --- 6. Coleta na Base OpenAlex ---
    add_heading_styled(doc, "6. Coleta na Base OpenAlex", 1)
    
    add_body_paragraph(doc,
        "O coletor do OpenAlex explora a flexibilidade da API global mantida pela organização sem fins lucrativos OurResearch. O OpenAlex atua como um agregador aberto e massivo de metadados acadêmicos de toda a internet, sendo extremamente rápido por não exigir chaves de acesso obrigatórias. O script envia requisições estruturadas contendo um User-Agent identificado com o e-mail do pesquisador. Isso inclui as requisições no pool de cortesia (Polite Pool) da API, garantindo limites de requisição ampliados e menor tempo de resposta nas consultas."
    )
    
    add_body_paragraph(doc,
        "Uma funcionalidade de grande relevância bibliométrica implementada no coletor do OpenAlex é a identificação unificada das fontes e repositórios originais de cada trabalho acadêmico. Como o OpenAlex opera como um agregador indireto, os registros brutos são compostos por uma árvore contendo múltiplos locais de hospedagem (locations). O script varre essa árvore de localizações e consolida as marcas das bases de dados reais (como PubMed, DOAJ, Dialnet ou SciELO) em uma coluna dedicada. Isso permite ao pesquisador identificar a sobreposição de indexação nas diferentes plataformas de pesquisa."
    )
    
    add_body_paragraph(doc,
        "Os resumos indexados na API do OpenAlex são armazenados no formato de índice invertido (abstract inverted index) por razões de direitos autorais de distribuição de textos. O coletor contém um algoritmo matemático eficiente que lê este dicionário de frequências de termos e reconstrói o texto do resumo original. A reconstrução posiciona cada palavra no seu respectivo índice numérico, restaurando o resumo legível para análises de mineração de texto posteriores. O design do Excel do OpenAlex utiliza tons elegantes de roxo-real."
    )
    
    # --- 7. Metodologia Híbrida de Enriquecimento ---
    add_heading_styled(doc, "7. Metodologia Híbrida de Enriquecimento de Resumos", 1)
    
    add_body_paragraph(doc,
        "Para contornar as severas restrições comerciais da Elsevier na recuperação de resumos, o sistema integra uma metodologia de enriquecimento híbrido nos coletores. Quando a busca na Scopus é concluída, o script inicia uma rotina que analisa a disponibilidade do resumo e das palavras-chave de cada item. Se o resumo estiver indisponível devido a restrições de licença da chave de API utilizada, o enriquecedor entra em ação. O algoritmo utiliza o DOI do artigo da Scopus para fazer uma consulta pontual direta na API gratuita do OpenAlex."
    )
    
    add_body_paragraph(doc,
        "Caso o artigo da Scopus não possua um DOI cadastrado na base de dados, a rotina realiza uma busca secundária no OpenAlex baseada no título textual limpo da obra. Esse método de busca por correspondência exata de título ou identificador evita erros comuns de falsos positivos na importação de textos, empregando mecanismos de requisição e scraping polido recomendados por especialistas (CHAPAGAIN, 2023). A combinação dessas duas fontes de dados garante planilhas robustas com elevado preenchimento de resumos e palavras-chave bibliométricas."
    )
    
    # --- 8. Desenho de Pesquisa Autoguiado e Gestão no Excel ---
    add_heading_styled(doc, "8. Formulário de Desenho de Pesquisa Autoguiado e Gestão no Excel", 1)
    
    add_body_paragraph(doc,
        "Para unificar o ciclo de vida da revisão sistemática sob as melhores práticas metodológicas do Cochrane Handbook (2019) e do manual Campbell (2020), o sistema incorpora um Formulário de Desenho de Pesquisa em formato Excel (Formulario_Desenho_Pesquisa_RSL.xlsx). Este arquivo atua como um guia estruturado composto por 14 abas organizadas em duas grandes fases. A fase de desenho orienta a equipe de pesquisa nas definições de identificação do protocolo, pergunta de pesquisa (framework PICO/PICo), critérios de elegibilidade temporal e idiomática, montagem dos blocos conceituais booleanos de busca, seleção de fontes de dados e definição das variáveis de extração e ferramentas de risco de viés (ROB 2, NOS, JBI)."
    )
    
    add_body_paragraph(doc,
        "A segunda fase do arquivo Excel compreende as planilhas operacionais e de gestão. A 'Planilha Triagem' habilita a condução do screening em dupla independente para as etapas de título/resumo e texto completo, contendo controles de indicação de conflito, motivos de exclusão e marcação de estágio do fluxo PRISMA. A 'Planilha Extração' e a 'Planilha Qualidade' são populadas dinamicamente com as variáveis do protocolo e fornecem recursos visuais como validações de dados por dropdown e formatação condicional automática para o risco de viés (sinalizando em verde baixo risco, amarelo preocupação e vermelho alto risco). A aba 'PRISMA Flow' consolida a contagem final de registros em tempo real para alimentar o diagrama PRISMA 2020."
    )
    
    add_body_paragraph(doc,
        "A integração entre o planejamento em Excel e o motor de coleta em Python é realizada de forma automatizada por meio do script de exportação (exportar_configs_do_excel.py). O script analisa sintaticamente a planilha Excel, localizando dinamicamente as expressões de busca montadas, os limites temporais, o e-mail cadastrado e as credenciais de API. Em seguida, gera de forma limpa os arquivos JSON de configuração (config.json, config_scopus.json, config_openalex.json e config_scielo.json) que alimentam os harvesters, eliminando a necessidade de edição manual de código ou JSONs pelo pesquisador."
    )
    
    # --- 9. Pipeline de Consolidação, Deduplicação e Análise Estatística ---
    add_heading_styled(doc, "9. Pipeline de Consolidação, Deduplicação e Análise Estatística", 1)
    
    add_body_paragraph(doc,
        "A fase de execução pós-coleta é sustentada por três ferramentas automatizadas que operacionalizam a triagem, extração de dados e a síntese final. O script de consolidação (consolidar_triagem.py) lê os arquivos CSV de saída de todos os coletores e unifica os registros em um único formato tabular. Em seguida, executa uma deduplicação em duas fases: primeiro, por correspondência exata do DOI normalizado; segundo, por similaridade difusa baseada no algoritmo SequenceMatcher (com limite de aceitação de 90%) nos títulos para registros que não contêm o DOI. O script também realiza uma pré-triagem automática eliminando trabalhos fora do recorte temporal ou em idiomas não elegíveis (permitindo por padrão inglês, português e espanhol), popula a aba 'Planilha Triagem' do formulário Excel e atualiza a aba 'PRISMA Flow' com os números do fluxo inicial."
    )
    
    add_body_paragraph(doc,
        "O processo de seleção final e sistematização das evidências é realizado pelo script de simulação e triagem (executar_triagem_e_extracao.py). Esta ferramenta simula decisões em dupla independente de revisão (Revisor A e Revisor B) tanto para a triagem por título e resumo (T/R) quanto para a triagem por texto completo (TC), preenchendo as justificativas e calculando os conflitos metodológicos diretamente na planilha Excel. Para os estudos aprovados, o script extrai automaticamente metadados estruturados (autor, ano, base, DOI, título, periódico, resumo, objetivos, desenho do estudo, tamanho da amostra, métodos estatísticos/algoritmos causais, limitações e dados geográficos) e os escreve na 'Planilha Extração', enquanto avalia os riscos de viés (ROB 2/ROBINS-I) na 'Planilha Qualidade', atualizando por fim os campos finais do fluxograma PRISMA."
    )
    
    add_body_paragraph(doc,
        "Por fim, o script de análise e síntese (gerar_sintese_final.py) processa os dados finais extraídos no Excel para realizar estatísticas descritivas automáticas do levantamento bibliométrico. A ferramenta compila a distribuição temporal da literatura, a tipologia dos desenhos de estudo, a frequência de uso de métodos de descoberta e inferência causal, os níveis de risco de viés avaliados, as limitações metodológicas reportadas e a distribuição geográfica das publicações. Ao término da análise, o script consolida os dados estruturados em tabelas formatadas em Markdown, gerando o relatório final completo de síntese de evidências (relatorio_sintese_final.md) para documentar as conclusões da revisão."
    )
    
    # --- 10. Como Operar o Sistema Localmente ---
    add_heading_styled(doc, "10. Como Operar o Sistema Localmente", 1)
    
    add_body_paragraph(doc,
        "A operação da suite de coletores foi otimizada para ser realizada em poucos passos utilizando a planilha Excel como fonte única de verdade. O pesquisador deve preencher os parâmetros do protocolo de pesquisa diretamente no Formulario_Desenho_Pesquisa_RSL.xlsx e salvar o arquivo. Em seguida, deve executar o script de exportação (exportar_configs_do_excel.py) no console para gravar as configurações. Com os arquivos de configuração gerados automaticamente, a coleta é disparada com um duplo clique no arquivo inicializador rápido correspondente (ex: run_scopus.bat ou run_openalex.bat) presente no diretório."
    )
    
    add_body_paragraph(doc,
        "O terminal exibirá em tempo real o andamento da paginação da busca, o status da coleta e as etapas de enriquecimento de resumos. Ao término do processamento, as pastas específicas de saída (ex: scielo_outputs, scopus_outputs e openalex_outputs) serão criadas e populadas no mesmo diretório. Em seguida, o pipeline de consolidação, triagem e síntese pode ser acionado sequencialmente (consolidar_triagem.py, executar_triagem_e_extracao.py e gerar_sintese_final.py) para gerar as planilhas consolidadas e o relatório de síntese narrativa. As planilhas geradas possuem recursos de formatação automática que ajustam a largura das colunas e habilitam a quebra de texto nas células."
    )
    
    # --- 11. Conclusão ---
    add_heading_styled(doc, "11. Conclusão", 1)
    
    add_body_paragraph(doc,
        "O desenvolvimento da suite de coletores científicos integrada ao desenho de pesquisa em Excel resolveu os gargalos críticos enfrentados na condução de revisões sistemáticas do projeto de pesquisa atual. A união entre o planejamento metodológico estruturado em Excel, o motor de enriquecimento híbrido de dados em Python, os scripts de deduplicação e triagem automatizados e as planilhas de gestão operacionais do protocolo proveu um fluxo de trabalho extremamente eficiente e reprodutível. A automação das configurações reduziu a barreira técnica para os pesquisadores, eliminando erros humanos de parametrização."
    )
    
    add_body_paragraph(doc,
        "Por fim, a padronização das saídas estruturadas em CSV, Excel e Markdown agilizou o processo de importação e triagem bibliográfica da literatura recuperada. Os coletores garantem a reprodutibilidade metodológica do levantamento bibliográfico, requisito essencial para a validação das revisões sistemáticas em periódicos de alto impacto científico. As ferramentas estão prontas para apoiar novos projetos e expressões de busca, automatizando de forma confiável o acesso à ciência mundial."
    )
    
    # --- 12. Referências Bibliográficas ---
    add_heading_styled(doc, "12. Referências Bibliográficas", 1)
    
    add_body_paragraph(doc,
        "CHAPAGAIN, Anish. Hands-On Web Scraping with Python: Extract quality data from the web using effective Python techniques. Birmingham: Packt Publishing, 2023. Este livro serviu de base conceitual para o desenvolvimento dos mecanismos de bypass de CDN e estruturação das requisições com headers realistas de navegadores comerciais, permitindo contornar firewalls de segurança como o Bunny Shield da SciELO de forma ética e estável durante a varredura."
    )
    
    add_body_paragraph(doc,
        "HIGGINS, J. P. T. et al. (eds). Cochrane Handbook for Systematic Reviews of Interventions version 6.0. Cochrane, 2019. O manual fornece a base de diretrizes para o desenho do protocolo de revisão sistemática e a formulação da questão PICO de pesquisa."
    )
    
    add_body_paragraph(doc,
        "IBICT. Padrão de Metadados da BDTD: MTD3-BR. Brasília: Instituto Brasileiro de Informação em Ciência e Tecnologia (IBICT), 2017. O manual técnico governamental estabelece as diretrizes de XML e os padrões de metadados para teses e dissertações brasileiras. Sua especificação orientou diretamente as rotinas de mapeamento de campos e a exportação final dos dados bibliográficos da BDTD no coletor específico deste projeto."
    )
    
    add_body_paragraph(doc,
        "MITCHELL, Ryan. Web Scraping with Python: Data Extraction from the Modern Web. 3. ed. Sebastopol: O'Reilly Media, 2024. A obra clássica detalha técnicas fundamentais de engenharia de dados, paginação e parseamento de HTML via BeautifulSoup. Suas orientações foram aplicadas diretamente na concepção lógica do coletor do portal SciELO e nas rotinas de auditoria, limpeza estrutural e normalização de textos dos resumos bibliográficos."
    )
    
    add_body_paragraph(doc,
        "PAGE, M. J. et al. The PRISMA 2020 statement. BMJ, v. 372, n. n71, 2021. Declaração que orienta a elaboração de fluxogramas de seleção de literatura e reporte transparente para revisões sistemáticas."
    )
    
    # Save the document in the workspace
    doc_path = "d:\\repos\\Revisão sistemática\\Metodologia_e_Operacao_Coletores.docx"
    doc.save(doc_path)
    print(f"Word document saved successfully at: {doc_path}")
    
if __name__ == "__main__":
    main()
