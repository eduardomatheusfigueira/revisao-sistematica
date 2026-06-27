import os
import re
import csv
import json
import time
import logging
from dataclasses import dataclass, asdict
from typing import List, Dict, Any, Optional
import requests
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import argparse

# Configure logging
def setup_logging(log_path: str):
    logger = logging.getLogger("BDTDHarvester")
    logger.setLevel(logging.DEBUG)
    
    if not logger.handlers:
        # File handler
        file_handler = logging.FileHandler(log_path, encoding="utf-8")
        file_handler.setLevel(logging.DEBUG)
        file_formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(name)s - %(message)s')
        file_handler.setFormatter(file_formatter)
        logger.addHandler(file_handler)
        
        # Console handler
        console_handler = logging.StreamHandler()
        console_handler.setLevel(logging.INFO)
        console_formatter = logging.Formatter('%(levelname)s: %(message)s')
        console_handler.setFormatter(console_formatter)
        logger.addHandler(console_handler)
    return logger

@dataclass
class BDTDRecord:
    id: str
    oai_id: str = "N/A"
    title: str = "N/A"
    alt_title: str = "N/A"
    authors: str = "N/A"
    advisors: str = "N/A"
    coadvisors: str = "N/A"
    referees: str = "N/A"
    year: int = 0
    date_issued: str = "N/A"
    format: str = "N/A"
    type: str = "N/A"
    institution_initials: str = "N/A"
    institution_fullname: str = "N/A"
    program: str = "N/A"
    department: str = "N/A"
    subjects: str = "N/A"
    subjects_cnpq: str = "N/A"
    language: str = "N/A"
    rights: str = "N/A"
    sponsorship: str = "N/A"
    abstract: str = "N/A"
    url: str = "N/A"

    @classmethod
    def from_solr_raw(cls, raw: Dict[str, Any]) -> 'BDTDRecord':
        """
        Maps Solr rawData keys (adhering to MTD3-BR) to BDTDRecord fields.
        """
        solr_raw = raw.get("rawData", {})
        
        # Helper to join multi-valued string lists
        def join_list(val_list: Any, separator: str = "; ") -> str:
            if not val_list:
                return "N/A"
            if isinstance(val_list, list):
                cleaned = [str(v).strip() for v in val_list if v and str(v).strip()]
                return separator.join(cleaned) if cleaned else "N/A"
            return str(val_list).strip() or "N/A"

        # Helper to get the first value of a multi-valued list or single string
        def first_val(val: Any) -> str:
            if not val:
                return "N/A"
            if isinstance(val, list):
                return str(val[0]).strip() if val[0] else "N/A"
            return str(val).strip() or "N/A"

        # Extract referees (banca) dynamically by checking keys dc.contributor.referee*
        referee_vals = []
        for k, v in sorted(solr_raw.items()):
            if "referee" in k.lower() and isinstance(v, list):
                for val in v:
                    if val and str(val).strip():
                        referee_vals.append(str(val).strip())
        referees_str = "; ".join(referee_vals) if referee_vals else "N/A"

        # Extract alternative title
        alt_title_str = join_list(solr_raw.get("dc.title.alternative.none.fl_str_mv"))
        
        # Extract advisors and coadvisors
        advisors_list = []
        for key in ["dc.contributor.advisor1.fl_str_mv", "dc.contributor.advisor2.fl_str_mv"]:
            val = solr_raw.get(key)
            if val:
                if isinstance(val, list):
                    advisors_list.extend([v for v in val if v])
                else:
                    advisors_list.append(val)
        if not advisors_list:
            advisors_list = solr_raw.get("contributor_str_mv", [])
        advisors_str = "; ".join(advisors_list) if advisors_list else "N/A"

        coadvisors_list = []
        for key in ["dc.contributor.advisor-co1.fl_str_mv", "dc.contributor.advisor-co2.fl_str_mv"]:
            val = solr_raw.get(key)
            if val:
                if isinstance(val, list):
                    coadvisors_list.extend([v for v in val if v])
                else:
                    coadvisors_list.append(val)
        coadvisors_str = "; ".join(coadvisors_list) if coadvisors_list else "N/A"

        # Extract subjects
        subjects_list = solr_raw.get("dc.subject.por.fl_str_mv", solr_raw.get("topic", []))
        subjects_str = join_list(subjects_list)
        subjects_cnpq_str = join_list(solr_raw.get("dc.subject.cnpq.fl_str_mv"))

        # Extract sponsorship
        sponsorship_str = join_list(solr_raw.get("dc.description.sponsorship.fl_str_mv"))

        # Extract abstract
        abstract_str = solr_raw.get("description")
        if not abstract_str:
            abstract_str = first_val(raw.get("summary"))
        else:
            abstract_str = str(abstract_str).strip()

        # Extract institution sigla
        inst_initials = first_val(solr_raw.get("network_acronym_str"))
        if inst_initials == "N/A":
            inst_initials = first_val(solr_raw.get("instacron_str"))
        if inst_initials == "N/A":
            inst_initials = first_val(solr_raw.get("institution"))

        # Clean fields
        rec = cls(
            id=raw.get("id", "N/A"),
            oai_id=solr_raw.get("oai_identifier_str", "N/A"),
            title=raw.get("title", "N/A"),
            alt_title=alt_title_str,
            authors=join_list(solr_raw.get("author", solr_raw.get("dc.contributor.author.fl_str_mv", []))),
            advisors=advisors_str,
            coadvisors=coadvisors_str,
            referees=referees_str,
            date_issued=first_val(solr_raw.get("dc.date.issued.fl_str_mv")),
            format=join_list(solr_raw.get("format")),
            type=first_val(solr_raw.get("dc.type.driver.fl_str_mv")),
            institution_initials=inst_initials,
            institution_fullname=first_val(solr_raw.get("dc.publisher.none.fl_str_mv")),
            program=first_val(solr_raw.get("dc.publisher.program.fl_str_mv")),
            department=first_val(solr_raw.get("dc.publisher.department.fl_str_mv")),
            subjects=subjects_str,
            subjects_cnpq=subjects_cnpq_str,
            language=join_list(solr_raw.get("language", solr_raw.get("dc.language.iso.fl_str_mv", []))),
            rights=join_list(solr_raw.get("dc.rights.driver.fl_str_mv")),
            sponsorship=sponsorship_str,
            abstract=abstract_str,
            url=first_val(solr_raw.get("url", solr_raw.get("dc.identifier.uri.fl_str_mv", [])))
        )
        return rec

class DataCleaner:
    @staticmethod
    def clean_text(text: str) -> str:
        if not text or text == "N/A":
            return "N/A"
        # Remove HTML tags
        text = re.sub(r'<[^<]+?>', '', text)
        # Normalize whitespace (replace tabs, newlines, multiple spaces with a single space)
        text = re.sub(r'\s+', ' ', text)
        return text.strip()

    @staticmethod
    def parse_year(date_str: str, publish_dates: Optional[List[str]] = None) -> int:
        if date_str and date_str != "N/A":
            match = re.search(r'\b(20\d{2}|19\d{2})\b', date_str)
            if match:
                return int(match.group(1))
        if publish_dates:
            for date in publish_dates:
                match = re.search(r'\b(20\d{2}|19\d{2})\b', str(date))
                if match:
                    return int(match.group(1))
        return 0

    @staticmethod
    def map_document_type(type_str: str, format_str: str) -> str:
        combined = f"{type_str} {format_str}".lower()
        if "disserta" in combined or "master" in combined:
            return "Dissertação"
        if "tese" in combined or "doctoral" in combined or "doctor" in combined:
            return "Tese"
        return "Tese/Dissertação"

    @classmethod
    def clean_record(cls, record: BDTDRecord, raw_pub_dates: Optional[List[str]] = None) -> BDTDRecord:
        record.title = cls.clean_text(record.title)
        record.alt_title = cls.clean_text(record.alt_title)
        record.authors = cls.clean_text(record.authors)
        record.advisors = cls.clean_text(record.advisors)
        record.coadvisors = cls.clean_text(record.coadvisors)
        record.referees = cls.clean_text(record.referees)
        record.institution_fullname = cls.clean_text(record.institution_fullname)
        record.program = cls.clean_text(record.program)
        record.department = cls.clean_text(record.department)
        record.subjects = cls.clean_text(record.subjects)
        record.subjects_cnpq = cls.clean_text(record.subjects_cnpq)
        record.sponsorship = cls.clean_text(record.sponsorship)
        record.abstract = cls.clean_text(record.abstract)
        
        # Parse year
        record.year = cls.parse_year(record.date_issued, raw_pub_dates)
        # Map type
        record.type = cls.map_document_type(record.type, record.format)
        
        return record

class BDTDClient:
    def __init__(self, base_url: str, user_agent: str, politeness_delay: float = 1.5, max_retries: int = 5, backoff_factor: float = 1.5):
        self.base_url = base_url
        self.politeness_delay = politeness_delay
        self.max_retries = max_retries
        self.backoff_factor = backoff_factor
        self.session = requests.Session()
        self.headers = {
            "User-Agent": user_agent,
            "Accept": "application/json"
        }
        self.logger = logging.getLogger("BDTDHarvester.Client")

    def search_page(self, query: str, page: int, limit: int = 50) -> Optional[Dict[str, Any]]:
        """
        Fetches search results from BDTD VuFind REST API with Exponential Backoff retry strategy.
        """
        params = {
            "lookfor": query,
            "type": "AllFields",
            "page": page,
            "limit": limit,
            "field[]": [
                "id", "title", "authors", "formats", "languages", 
                "subjects", "urls", "summary", "institutions", 
                "publicationDates", "rawData"
            ]
        }
        
        url = self.base_url
        self.logger.debug(f"Requesting API - Page: {page} | Query: {query}")
        
        for attempt in range(1, self.max_retries + 1):
            try:
                response = self.session.get(url, params=params, headers=self.headers, timeout=25)
                
                if response.status_code == 200:
                    return response.json()
                
                # Check rate limits or server errors
                if response.status_code == 429:
                    sleep_time = self.backoff_factor * (2 ** attempt)
                    self.logger.warning(f"Attempt {attempt}/{self.max_retries}: HTTP 429 (Too Many Requests). Waiting {sleep_time:.2f}s...")
                    time.sleep(sleep_time)
                elif response.status_code >= 500:
                    sleep_time = self.backoff_factor * (2 ** attempt)
                    self.logger.warning(f"Attempt {attempt}/{self.max_retries}: HTTP {response.status_code} (Server Error). Waiting {sleep_time:.2f}s...")
                    time.sleep(sleep_time)
                else:
                    self.logger.error(f"HTTP Error {response.status_code} encountered. No retry for this code.")
                    break
                    
            except (requests.exceptions.RequestException, requests.exceptions.Timeout) as e:
                sleep_time = self.backoff_factor * (2 ** attempt)
                self.logger.warning(f"Attempt {attempt}/{self.max_retries}: Network error ({e}). Waiting {sleep_time:.2f}s...")
                time.sleep(sleep_time)
                
        return None

class DataExporter:
    def __init__(self, output_dir: str, logger: logging.Logger):
        self.output_dir = output_dir
        self.logger = logger
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
            self.logger.info(f"Created output directory: {output_dir}")

    def export_to_json(self, records: List[BDTDRecord], filename: str):
        json_path = os.path.join(self.output_dir, filename)
        self.logger.info(f"Exporting JSON raw backup: {json_path}")
        dict_records = [asdict(r) for r in records]
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(dict_records, f, indent=2, ensure_ascii=False)

    def export_to_csv(self, records: List[BDTDRecord], filename: str):
        csv_path = os.path.join(self.output_dir, filename)
        self.logger.info(f"Exporting CSV: {csv_path}")
        
        headers = [
            "ID BDTD", "OAI ID", "Título do Trabalho", "Título Alternativo", "Autores", 
            "Orientadores", "Coorientadores", "Banca de Defesa", "Ano", 
            "Data de Defesa", "Formato", "Tipo de Trabalho", "Instituição (Sigla)", 
            "Instituição (Nome)", "Programa", "Departamento", "Palavras-chave", 
            "Área CNPq", "Idioma", "Direitos de Acesso", "Agência de Fomento", 
            "Link / URL", "Resumo"
        ]
        
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for r in records:
                writer.writerow([
                    r.id, r.oai_id, r.title, r.alt_title, r.authors,
                    r.advisors, r.coadvisors, r.referees, r.year,
                    r.date_issued, r.format, r.type, r.institution_initials,
                    r.institution_fullname, r.program, r.department, r.subjects,
                    r.subjects_cnpq, r.language, r.rights, r.sponsorship,
                    r.url, r.abstract
                ])

    def export_to_excel(self, records: List[BDTDRecord], filename: str):
        excel_path = os.path.join(self.output_dir, filename)
        self.logger.info(f"Creating a brand-new Excel workbook: {excel_path}")
        
        wb = openpyxl.Workbook()
        
        # 1. BDTD Cleaned Data Sheet
        ws_data = wb.active
        ws_data.title = "BDTD Cleaned Data"
        
        headers = [
            "ID BDTD", "OAI ID", "Título", "Autores", "Orientadores", "Coorientadores", 
            "Ano de Defesa", "Tipo", "Instituição (Sigla)", "Instituição (Nome)", 
            "Programa", "Palavras-chave", "Idioma", "Agência de Fomento", "Link / URL"
        ]
        ws_data.append(headers)
        
        # Styling Setup
        font_family = "Segoe UI"
        font_header = Font(name=font_family, size=11, bold=True, color="FFFFFF")
        font_regular = Font(name=font_family, size=10, color="333333")
        thin_side = Side(border_style="thin", color="D3D3D3")
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        fill_header = PatternFill(start_color="366092", end_color="366092", fill_type="solid") # Dark steel blue
        fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        fill_zebra = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        
        # Format Headers
        ws_data.row_dimensions[1].height = 28
        for col_idx in range(1, len(headers) + 1):
            cell = ws_data.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_all
            
        # Append and format data
        for row_idx, r in enumerate(records, 2):
            ws_data.row_dimensions[row_idx].height = 24
            row_data = [
                r.id, r.oai_id, r.title, r.authors, r.advisors, r.coadvisors,
                r.year, r.type, r.institution_initials, r.institution_fullname,
                r.program, r.subjects, r.language, r.sponsorship, r.url
            ]
            fill = fill_zebra if row_idx % 2 == 0 else fill_white
            for col_idx, val in enumerate(row_data, 1):
                cell = ws_data.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_regular
                cell.border = border_all
                cell.fill = fill
                if col_idx in [1, 2, 7, 8, 9, 13]:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left
        
        # 2. Summary Sheet
        ws_summary = wb.create_sheet("Extraction Summary")
        ws_summary.row_dimensions[2].height = 26
        ws_summary.cell(row=2, column=2, value="BDTD Extraction Summary").font = Font(name=font_family, size=16, bold=True, color="366092")
        
        # Date & Count Info
        ws_summary.cell(row=4, column=2, value="Execution Date:").font = Font(name=font_family, size=10, bold=True)
        ws_summary.cell(row=4, column=3, value=time.strftime('%Y-%m-%d %H:%M:%S')).font = font_regular
        ws_summary.cell(row=5, column=2, value="Total Unique Works:").font = Font(name=font_family, size=10, bold=True)
        ws_summary.cell(row=5, column=3, value=len(records)).font = font_regular
        
        # Compilation stats for Document Types
        from collections import Counter
        types_count = Counter([r.type for r in records])
        
        ws_summary.cell(row=7, column=2, value="Document Type Summary").font = Font(name=font_family, size=12, bold=True, color="366092")
        ws_summary.cell(row=8, column=2, value="Document Type").font = font_header
        ws_summary.cell(row=8, column=2).fill = fill_header
        ws_summary.cell(row=8, column=3, value="Count").font = font_header
        ws_summary.cell(row=8, column=3).fill = fill_header
        
        cur_row = 9
        for t, count in types_count.items():
            ws_summary.row_dimensions[cur_row].height = 20
            c_type = ws_summary.cell(row=cur_row, column=2, value=t)
            c_count = ws_summary.cell(row=cur_row, column=3, value=count)
            c_type.font = font_regular
            c_count.font = font_regular
            c_type.border = border_all
            c_count.border = border_all
            c_type.alignment = align_left
            c_count.alignment = align_center
            cur_row += 1
            
        wb.save(excel_path)
        self.logger.info(f"Excel workbook created and saved successfully at: {excel_path}")

    def generate_analysis_report(self, records: List[BDTDRecord], filename: str, query: str):
        report_path = os.path.join(self.output_dir, filename)
        self.logger.info(f"Generating markdown summary report: {report_path}")
        
        from collections import Counter
        
        total = len(records)
        years = [r.year for r in records if r.year > 0]
        types = [r.type for r in records]
        institutions = [r.institution_initials for r in records if r.institution_initials != "N/A"]
        languages = [r.language for r in records if r.language != "N/A"]
        
        years_count = Counter(years)
        types_count = Counter(types)
        inst_count = Counter(institutions)
        lang_count = Counter(languages)
        
        keywords_list = []
        for r in records:
            if r.subjects and r.subjects != "N/A":
                keywords_list.extend([kw.strip().lower() for kw in r.subjects.split(";") if kw.strip()])
        kw_count = Counter(keywords_list)
        
        advisors_list = []
        for r in records:
            if r.advisors and r.advisors != "N/A":
                advisors_list.extend([adv.strip() for adv in r.advisors.split(";") if adv.strip()])
        adv_count = Counter(advisors_list)

        report = []
        report.append("# BDTD Scientific Data Extraction Report")
        report.append(f"**Coleta realizada em**: {time.strftime('%d/%m/%Y %H:%M:%S')}")
        report.append(f"**Expressão de Busca**: `{query}`")
        report.append(f"**Total de trabalhos exclusivos recuperados**: {total}\n")
        
        report.append("## 1. Cronologia de Publicações (Defesa)")
        report.append("| Ano | Quantidade | Percentual |")
        report.append("| --- | --- | --- |")
        for yr, count in sorted(years_count.items(), reverse=True):
            pct = (count / total) * 100
            report.append(f"| {yr} | {count} | {pct:.1f}% |")
        if not years:
            report.append("| N/A | 0 | 0.0% |")
        report.append("")
        
        report.append("## 2. Tipologia Documental")
        report.append("| Tipo de Trabalho | Quantidade | Percentual |")
        report.append("| --- | --- | --- |")
        for tp, count in types_count.most_common():
            pct = (count / total) * 100
            report.append(f"| {tp} | {count} | {pct:.1f}% |")
        report.append("")
        
        report.append("## 3. Distribuição Geográfica Institucional (Top 10)")
        report.append("| Instituição (Sigla) | Trabalhos | Percentual |")
        report.append("| --- | --- | --- |")
        for inst, count in inst_count.most_common(10):
            pct = (count / total) * 100
            report.append(f"| {inst} | {count} | {pct:.1f}% |")
        report.append("")
        
        report.append("## 4. Distribuição por Idiomas")
        report.append("| Idioma | Trabalhos | Percentual |")
        report.append("| --- | --- | --- |")
        for lang, count in lang_count.most_common():
            pct = (count / total) * 100
            report.append(f"| {lang} | {count} | {pct:.1f}% |")
        report.append("")

        report.append("## 5. Orientadores mais Frequentes (Top 10)")
        report.append("| Orientador | Trabalhos |")
        report.append("| --- | --- |")
        for adv, count in adv_count.most_common(10):
            report.append(f"| {adv} | {count} |")
        report.append("")

        report.append("## 6. Principais Palavras-chave (Top 15)")
        report.append("| Termo de Assunto | Frequência |")
        report.append("| --- | --- |")
        for kw, count in kw_count.most_common(15):
            report.append(f"| {kw} | {count} |")
        report.append("")
        
        with open(report_path, "w", encoding="utf-8") as f:
            f.write("\n".join(report))

def main():
    # Setup argparse CLI
    parser = argparse.ArgumentParser(description="BDTD Scientific Data Harvester - CLI Tool")
    parser.add_argument("--query", type=str, help="Search query expression for BDTD (e.g. 'saneamento AND \"inferência causal\"')")
    parser.add_argument("--start-year", type=int, help="Start year constraint (e.g. 2000)")
    parser.add_argument("--end-year", type=int, help="End year constraint (e.g. 2026)")
    parser.add_argument("--output-dir", type=str, help="Directory to save output files")
    parser.add_argument("--limit", type=int, help="Number of records per page (max 100)")
    parser.add_argument("--config", type=str, default="config.json", help="Path to config JSON file")
    args = parser.parse_args()

    # Load configuration if present
    config = {}
    if os.path.exists(args.config):
        try:
            with open(args.config, "r", encoding="utf-8") as f:
                config = json.load(f)
        except Exception as e:
            print(f"Warning: Failed to load config file '{args.config}': {e}. Using CLI arguments/defaults.")

    # Resolve settings (CLI has highest priority, then config file, then hardcoded defaults)
    search_cfg = config.get("search", {})
    query = args.query or search_cfg.get("query")
    
    if not query:
        print("\n[ERROR] No search query provided! Please configure it in config.json or use the --query argument.")
        parser.print_help()
        return

    start_year = args.start_year or search_cfg.get("start_year") or 2000
    end_year = args.end_year or search_cfg.get("end_year") or 2026

    paths = config.get("paths", {})
    output_dir = args.output_dir or paths.get("output_dir") or "data_outputs"
    log_name = paths.get("log_name") or "bdtd_harvester.log"
    csv_name = paths.get("csv_name") or "bdtd_clean_data.csv"
    json_name = paths.get("json_name") or "bdtd_raw_backup.json"
    excel_name = paths.get("excel_name") or "BDTD_Data_Export.xlsx"
    report_name = paths.get("report_name") or "bdtd_summary_report.md"

    # Setup directories and logger
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
        
    logger = setup_logging(os.path.join(output_dir, log_name))
    logger.info("=== BDTD SYSTEM DATA HARVESTER STARTED (CLI MODE) ===")
    logger.info(f"Target query: '{query}'")
    logger.info(f"Temporal window: {start_year} - {end_year}")
    logger.info(f"Saving outputs to: {output_dir}")
    
    api_cfg = config.get("api", {})
    limit = args.limit or api_cfg.get("limit") or 50
    base_url = api_cfg.get("base_url") or "https://bdtd.ibict.br/vufind/api/v1/search"
    user_agent = api_cfg.get("user_agent") or "BDTDHarvester/1.0"
    politeness_delay = api_cfg.get("politeness_delay_seconds") or 1.5
    max_retries = api_cfg.get("max_retries") or 5
    backoff_factor = api_cfg.get("backoff_factor") or 1.5

    client = BDTDClient(
        base_url=base_url,
        user_agent=user_agent,
        politeness_delay=politeness_delay,
        max_retries=max_retries,
        backoff_factor=backoff_factor
    )
    
    raw_records: Dict[str, Dict[str, Any]] = {}
    page = 1
    
    logger.info(f"Initiating retrieval loop for query: '{query}'...")
    while True:
        logger.info(f"Requesting Page {page}...")
        response_data = client.search_page(query, page, limit)
        
        if not response_data:
            logger.error(f"Error fetching API Page {page}. Breaking query loop.")
            break
            
        records_list = response_data.get("records", [])
        result_count = response_data.get("resultCount", 0)
        
        if page == 1:
            logger.info(f"BDTD total matching records in index: {result_count}")
            
        if not records_list:
            logger.info("No more records found in this page. Harvesting complete.")
            break
            
        logger.info(f"Fetched {len(records_list)} records from page {page}.")
        
        for item in records_list:
            rec_id = item.get("id")
            if rec_id:
                raw_records[rec_id] = item
                
        # Break condition
        if len(records_list) < limit or (page * limit) >= result_count:
            logger.info("Finished harvesting all matches.")
            break
            
        page += 1
        time.sleep(client.politeness_delay)
        
    logger.info(f"Total unique records harvested: {len(raw_records)}")
    
    # Process, validate and clean data
    cleaned_records: List[BDTDRecord] = []
    for rec_id, raw_item in raw_records.items():
        try:
            record = BDTDRecord.from_solr_raw(raw_item)
            raw_pub_dates = raw_item.get("publicationDates", [])
            cleaned_record = DataCleaner.clean_record(record, raw_pub_dates)
            
            # Year constraint check
            if start_year <= cleaned_record.year <= end_year:
                cleaned_records.append(cleaned_record)
            else:
                logger.debug(f"Record {rec_id} filtered out by year ({cleaned_record.year}).")
        except Exception as e:
            logger.error(f"Error processing record {rec_id}: {e}", exc_info=True)
            
    logger.info(f"Processing finished. {len(cleaned_records)} records passed validation filters.")
    
    # Export Data
    exporter = DataExporter(output_dir, logger)
    
    # JSON raw backup
    exporter.export_to_json([BDTDRecord.from_solr_raw(item) for item in raw_records.values()], json_name)
    
    # CSV clean dataset
    exporter.export_to_csv(cleaned_records, csv_name)
    
    # Excel new workbook
    exporter.export_to_excel(cleaned_records, excel_name)
    
    # Markdown analysis
    exporter.generate_analysis_report(cleaned_records, report_name, query)
    
    logger.info("=== BDTD SYSTEM DATA HARVESTER PROCESS COMPLETED SUCCESSFULLY ===")

if __name__ == "__main__":
    main()
