"""
Script de Exportação de Configurações do Excel para os Coletores
==============================================================================
Lê o arquivo 'Formulario_Desenho_Pesquisa_RSL.xlsx' editado pelo pesquisador,
extrai as configurações de busca, recorte temporal e chaves de API, e gera/atualiza
os respectivos arquivos JSON de configuração:
  - config.json (BDTD)
  - config_scopus.json (Scopus)
  - config_openalex.json (OpenAlex)
  - config_scielo.json (SciELO)

Uso:
  python exportar_configs_do_excel.py
"""

import os
import json
import openpyxl

def find_value_by_label(ws, label_text, label_col="B", val_col="D"):
    """Varre as linhas da coluna especificada procurando o rótulo e retorna o valor correspondente."""
    for r in range(1, ws.max_row + 1):
        cell_val = ws[f"{label_col}{r}"].value
        if cell_val and label_text.lower() in str(cell_val).lower():
            # Retorna o valor da célula correspondente de input
            return ws[f"{val_col}{r}"].value
    return None

def get_selected_bases(ws):
    """Varre a tabela de bases na aba '5. Fontes de Dados' e retorna quais estão marcadas para usar."""
    selected = {}
    # A tabela de bases começa depois do cabeçalho da seção 5.1
    # Coluna C: Base, Coluna E: Usar? (S/N)
    for r in range(1, ws.max_row + 1):
        base_name = ws[f"C{r}"].value
        if base_name in ["BDTD", "SciELO", "Scopus", "OpenAlex"]:
            use_val = ws[f"E{r}"].value
            selected[base_name] = str(use_val).strip().lower() in ["sim", "s", "yes", "y", "1"]
    return selected

def main():
    print("=" * 60)
    print("  EXPORTADOR DE CONFIGURAÇÕES PARA OS COLETORES")
    print("=" * 60)
    print()
    
    excel_path = "Formulario_Desenho_Pesquisa_RSL.xlsx"
    if not os.path.exists(excel_path):
        print(f"[ERRO] O arquivo '{excel_path}' não foi encontrado.")
        print("Por favor, gere o arquivo primeiro executando: python gerar_formulario_excel.py")
        return

    print(f"Lendo '{excel_path}'...")
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
    except Exception as e:
        print(f"[ERRO] Falha ao abrir o Excel: {e}")
        return

    # 1. Aba de Estratégia de Busca
    if "4. Estratégia Busca" not in wb.sheetnames:
        print("[ERRO] Aba '4. Estratégia Busca' não encontrada no arquivo Excel.")
        return
    ws_search = wb["4. Estratégia Busca"]
    
    query = find_value_by_label(ws_search, "Expressão final *")
    start_year = find_value_by_label(ws_search, "Ano inicial")
    end_year = find_value_by_label(ws_search, "Ano final")
    scopus_field = find_value_by_label(ws_search, "Campo Scopus")
    
    # 2. Aba de Fontes de Dados
    if "5. Fontes de Dados" not in wb.sheetnames:
        print("[ERRO] Aba '5. Fontes de Dados' não encontrada no arquivo Excel.")
        return
    ws_sources = wb["5. Fontes de Dados"]
    
    scopus_key = find_value_by_label(ws_sources, "Scopus API Key")
    email = find_value_by_label(ws_sources, "E-mail Polite Pool")
    selected_bases = get_selected_bases(ws_sources)

    # Normalizar valores
    if not query:
        query = "\"inferência causal\""
        print("[AVISO] Expressão de busca não preenchida no Excel. Usando valor padrão.")
    
    try:
        start_year = int(start_year) if start_year else 2000
    except ValueError:
        start_year = 2000
        
    try:
        end_year = int(end_year) if end_year else 2026
    except ValueError:
        end_year = 2026
        
    scopus_field = str(scopus_field).strip() if scopus_field else "TITLE-ABS-KEY"
    scopus_key = str(scopus_key).strip() if scopus_key else "33698870c47d2706e3a3fc4c03397832"
    email = str(email).strip() if email else "user@mail.com"

    print("\nConfigurações extraídas do Excel:")
    print(f"  - Query: {query}")
    print(f"  - Recorte: {start_year} - {end_year}")
    print(f"  - Campo Scopus: {scopus_field}")
    print(f"  - Scopus Key: {'*** (configurada)' if scopus_key else 'Não configurada'}")
    print(f"  - E-mail Polite Pool: {email}")
    print("  - Bases selecionadas:")
    for base, is_sel in selected_bases.items():
        print(f"    * {base}: {'Sim' if is_sel else 'Não'}")

    print("\nGerando/Atualizando arquivos de configuração JSON...")

    # BDTD
    if selected_bases.get("BDTD", True):
        config_bdtd = {
            "search": { "query": query, "start_year": start_year, "end_year": end_year },
            "api": {
                "base_url": "https://bdtd.ibict.br/vufind/api/v1/search",
                "limit": 50, "max_retries": 5, "backoff_factor": 1.5, "politeness_delay_seconds": 1.5,
                "user_agent": f"BDTDHarvester/1.0 (contact: {email})"
            },
            "paths": {
                "output_dir": "data_outputs", "csv_name": "bdtd_clean_data.csv",
                "json_name": "bdtd_raw_backup.json", "excel_name": "BDTD_Data_Export.xlsx",
                "report_name": "bdtd_summary_report.md", "log_name": "bdtd_harvester.log"
            }
        }
        with open("config.json", "w", encoding="utf-8") as f:
            json.dump(config_bdtd, f, indent=2, ensure_ascii=False)
        print("  [OK] config.json (BDTD) atualizado.")

    # Scopus
    if selected_bases.get("Scopus", True):
        scopus_query = f"{scopus_field}({query})"
        config_scopus = {
            "search": { "query": scopus_query, "start_year": start_year, "end_year": end_year },
            "api": {
                "base_url": "https://api.elsevier.com/content/search/scopus",
                "api_key": scopus_key,
                "limit": 25, "max_retries": 5, "backoff_factor": 1.5, "politeness_delay_seconds": 1.0,
                "user_agent": f"ScopusHarvester/1.0 (contact: {email})"
            },
            "paths": {
                "output_dir": "scopus_outputs", "csv_name": "scopus_clean_data.csv",
                "json_name": "scopus_raw_backup.json", "excel_name": "Scopus_Data_Export.xlsx",
                "report_name": "scopus_summary_report.md", "log_name": "scopus_harvester.log"
            }
        }
        with open("config_scopus.json", "w", encoding="utf-8") as f:
            json.dump(config_scopus, f, indent=2, ensure_ascii=False)
        print("  [OK] config_scopus.json atualizado.")

    # OpenAlex
    if selected_bases.get("OpenAlex", True):
        config_openalex = {
            "search": {
                "query": query, "start_year": start_year, "end_year": end_year,
                "filters": { "repository_ids": [], "publisher_ids": [], "only_open_access": False, "source_types": [] }
            },
            "api": {
                "base_url": "https://api.openalex.org/works",
                "limit": 50, "max_retries": 5, "backoff_factor": 1.5, "politeness_delay_seconds": 1.0,
                "user_agent": f"OpenAlexHarvester/1.0 (contact: {email})"
            },
            "paths": {
                "output_dir": "openalex_outputs", "csv_name": "openalex_clean_data.csv",
                "json_name": "openalex_raw_backup.json", "excel_name": "OpenAlex_Data_Export.xlsx",
                "report_name": "openalex_summary_report.md", "log_name": "openalex_harvester.log"
            }
        }
        with open("config_openalex.json", "w", encoding="utf-8") as f:
            json.dump(config_openalex, f, indent=2, ensure_ascii=False)
        print("  [OK] config_openalex.json atualizado.")

    # SciELO
    if selected_bases.get("SciELO", True):
        config_scielo = {
            "search": { "query": query, "start_year": start_year, "end_year": end_year },
            "api": {
                "base_url": "https://api.openalex.org/works",
                "publisher_id": "P4310312277",
                "limit": 50, "max_retries": 5, "backoff_factor": 1.5, "politeness_delay_seconds": 1.0,
                "user_agent": f"SciELOHarvester/1.0 (contact: {email})"
            },
            "paths": {
                "output_dir": "scielo_outputs", "csv_name": "scielo_clean_data.csv",
                "json_name": "scielo_raw_backup.json", "excel_name": "SciELO_Data_Export.xlsx",
                "report_name": "scielo_summary_report.md", "log_name": "scielo_harvester.log"
            }
        }
        with open("config_scielo.json", "w", encoding="utf-8") as f:
            json.dump(config_scielo, f, indent=2, ensure_ascii=False)
        print("  [OK] config_scielo.json atualizado.")

    print("\n[SUCESSO] Processo de exportação concluído!")
    print("Agora você já pode executar os scripts de coleta correspondentes (ex: run_scopus.bat).")
    print("=" * 60)

if __name__ == "__main__":
    main()
